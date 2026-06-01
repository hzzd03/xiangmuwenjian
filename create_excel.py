from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

wb = Workbook()
ws = wb.active
ws.title = "自查表"

# Column widths
col_widths = {'A': 8, 'B': 28, 'C': 16, 'D': 16, 'E': 12, 'F': 14, 'G': 14,
              'H': 14, 'I': 14, 'J': 14, 'K': 16, 'L': 16, 'M': 18, 'N': 18, 'O': 12}
for col, w in col_widths.items():
    ws.column_dimensions[col].width = w

# Styles
title_font = Font(name='宋体', size=16, bold=True)
sub_font = Font(name='宋体', size=10, bold=True)
header_font = Font(name='宋体', size=9, bold=True)
data_font = Font(name='宋体', size=9)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
thin = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

# Row 1: 附件2
ws.merge_cells('A1:O1')
c = ws['A1']
c.value = '附件2'
c.font = Font(name='宋体', size=11, bold=True)
c.alignment = Alignment(horizontal='left', vertical='center')

# Row 2: Title
ws.merge_cells('A2:O2')
ws['A2'].value = '个人信息保护系列专项行动自查表'
ws['A2'].font = title_font
ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 38

# Row 3: 单位名称
ws.merge_cells('A3:O3')
ws['A3'].value = '单位名称（公章）：'
ws['A3'].font = sub_font
ws['A3'].alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[3].height = 32

# ---- Header Row 4: Top-level categories ----
h4 = [
    ('A4','B4','种类'),
    ('C4','C4','是否在无关场景\n收集个人信息'),
    ('D4','D4','是否向第三方\n机构提供\n个人信息'),
    ('E4','E4','是否泄露'),
    ('F4','G4','是否有其他\n违法违规行为'),
    ('H4','J4','收集使用个人信息'),
    ('K4','K4','个人信息主体\n权利保障'),
    ('L4','L4','是否建立\n个人信息保护\n管理制度'),
    ('M4','M4','是否对个人信息\n采取有效安全\n保护措施'),
    ('N4','N4','是否与运维单位\n专门签订\n安全保密协议'),
    ('O4','O4','备注'),
]
for s, e, txt in h4:
    if s != e:
        ws.merge_cells(f'{s}:{e}')
    c = ws[s]
    c.value = txt
    c.font = header_font
    c.alignment = center
    c.border = thin
    c.fill = fill
for col in range(1, 16):
    ws.cell(row=4, column=col).border = thin
    ws.cell(row=4, column=col).fill = fill
ws.row_dimensions[4].height = 50

# ---- Header Row 5: Sub-headers ----
h5 = [
    ('A5','A5','序号'),
    ('B5','B5','名称\n（APP/小程序/\n快应用/公众号/SDK等）'),
    ('C5','C5','是否在无关场景\n收集个人信息'),
    ('D5','D5','是否向第三方\n机构提供'),
    ('E5','E5','是否泄露'),
    ('F5','F5','是否强制收集'),
    ('G5','G5','是否告知\n收集使用规则'),
    ('H5','H5','是否经用户同意'),
    ('I5','I5','是否收集\n非必要个人信息'),
    ('J5','J5','通讯录等\n个人信息'),
    ('K5','K5','告知机构名称/\n地址、方式'),
    ('L5','L5','是否建立\n管理制度'),
    ('M5','M5','是否采取\n有效安全措施'),
    ('N5','N5','是否签订\n保密协议'),
    ('O5','O5','备注'),
]
for s, e, txt in h5:
    if s != e:
        ws.merge_cells(f'{s}:{e}')
    c = ws[s]
    c.value = txt
    c.font = header_font
    c.alignment = center
    c.border = thin
    c.fill = fill
ws.row_dimensions[5].height = 54

# ---- Data rows ----
for i in range(1, 6):
    row = 5 + i
    ws.cell(row=row, column=1).value = i
    for col in range(1, 16):
        c = ws.cell(row=row, column=col)
        c.font = data_font
        c.alignment = center
        c.border = thin
    ws.row_dimensions[row].height = 30

# ---- Footer ----
r = 11
ws.merge_cells(f'A{r}:H{r}')
ws[f'A{r}'].value = '填报人：'
ws[f'A{r}'].font = data_font
ws[f'A{r}'].alignment = Alignment(horizontal='left', vertical='center')
ws.merge_cells(f'I{r}:O{r}')
ws[f'I{r}'].value = '部门负责人：'
ws[f'I{r}'].font = data_font
ws[f'I{r}'].alignment = Alignment(horizontal='left', vertical='center')
ws.row_dimensions[r].height = 30

# Page setup
ws.page_setup.orientation = 'landscape'
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1

output_path = 'd:/xiangmuwenjian/assets/个人信息保护系列专项行动自查表.xlsx'
wb.save(output_path)
print(f'Excel saved to: {output_path}')
